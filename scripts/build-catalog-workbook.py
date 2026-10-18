"""Generate Excel-editable catalog templates and flat CSVs. Requires openpyxl for artifact generation only."""
from pathlib import Path
import csv
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

root=Path(__file__).resolve().parent.parent
out=root/'docs/catalog-starter'
base=['kind','externalKey','category','productUnit','verification','origin','sourceUrl','verificationBasis']
attributes=['diameterIn','material','type','schedule','lengthIn','endConnection','minSpanIn','maxSpanIn','mounting','nominalSpacing','widthIn','heightIn','depthIn','knockouts','specification']
columns={
 'material':base+attributes,
 'product':['kind','materialKey','manufacturer','partNumber','title','sourceUrl','verification','origin'],
 'translation':['kind','materialKey','phrase','origin','sourceUrl'],
 'image':['kind','materialKey','manufacturer','partNumber','sourceUrl','rightsBasis'],
}
examples={
 'material':dict(kind='material',externalKey='EXAMPLE-SUPPORT-001',category='bracket',productUnit='EA',verification='provisional',origin='ai_import',sourceUrl='https://example.invalid/fictional-only',verificationBasis='FICTIONAL SCHEMA EXAMPLE ONLY. Not a real product; do not load into the live catalog.',minSpanIn=12,maxSpanIn=20,mounting='FICTIONAL TEST MOUNT',material='FICTIONAL TEST MATERIAL',nominalSpacing='EXAMPLE ONLY'),
 'product':dict(kind='product',materialKey='EXAMPLE-SUPPORT-001',manufacturer='FICTIONAL EXAMPLE MAKER',partNumber='EXAMPLE-NOT-A-REAL-PART',title='FICTIONAL SCHEMA EXAMPLE ONLY',sourceUrl='https://example.invalid/fictional-only',verification='provisional',origin='ai_import'),
 'translation':dict(kind='translation',materialKey='EXAMPLE-SUPPORT-001',phrase='fictional example support bracket',origin='ai_generated',sourceUrl='https://example.invalid/fictional-only'),
 'image':dict(kind='image',materialKey='EXAMPLE-SUPPORT-001',manufacturer='FICTIONAL EXAMPLE MAKER',partNumber='EXAMPLE-NOT-A-REAL-PART',sourceUrl='https://example.invalid/fictional-only.png',rightsBasis='FICTIONAL PLACEHOLDER ONLY. No usage permission exists for an actual image.'),
}
for kind,headers in columns.items():
 for suffix,rows in [('template',[]),('FICTIONAL-EXAMPLE',[examples[kind]])]:
  with (out/f'{kind}-{suffix}.csv').open('w',encoding='utf-8-sig',newline='') as f:
   writer=csv.DictWriter(f,fieldnames=headers,lineterminator="\n");writer.writeheader();writer.writerows(rows)
sheets=[
 ('Conduit','material','conduit',['diameterIn','material','type','schedule','lengthIn','endConnection']),
 ('Supports','material','bracket',['minSpanIn','maxSpanIn','mounting','material','nominalSpacing']),
 ('Boxes','material','box',['type','widthIn','heightIn','depthIn','material','knockouts']),
 ('Other Parts','material','general',['type','specification']),
 ('Manufacturer Products','product',None,[]),('Translations','translation',None,[]),('Image References','image',None,[]),
]
numeric={'diameterIn','lengthIn','minSpanIn','maxSpanIn','widthIn','heightIn','depthIn'}
for fictional in [False,True]:
 wb=Workbook();guide=wb.active;guide.title='Instructions'
 notes=[['EtalonOps catalog workbook','FICTIONAL EXAMPLES — do not import as real stock' if fictional else 'Blank Excel template — enter your real source-backed parts'],
 ['How to edit','One record per row. No JSON or arrays. Keep header names unchanged.'],
 ['Catalog items','Use the appropriate category sheet. One canonical material key per meaningful specification, independent of brand.'],
 ['Manufacturer Products','One manufacturer part per row, linked using materialKey. Part numbers and keys are formatted as TEXT to preserve leading zeros.'],
 ['Translations','One search phrase per row, linked using materialKey. Repeat the material key for additional phrases.'],
 ['Image References','Optional source references; importing a row does not download or approve a photo.'],
 ['kind and category','Fill kind on every data row: material, product, translation or image. Material categories: conduit, bracket, box or general.'],
 ['Dimensions','Enter decimal inches as numbers: 15 3/4 inches = 15.75, 10 feet = 120. Keep actual span separate from nominalSpacing.'],
 ['AI provenance','AI-prepared materials/products: origin=ai_import, verification=provisional. Generated phrases: origin=ai_generated with sourceUrl.'],
 ['Description','The application builds the catalog description from category, specifications and productUnit. Manufacturer names belong in Manufacturer Products.'],
 ['Import','Save each populated data worksheet as CSV UTF-8 (Comma delimited). Upload CSV files, not this XLSX, to Administration → Catalog and imports.'],
 ['Order','Import all material/category files first, then Manufacturer Products, Translations and optional Image References.'],
 ['Limits','Each CSV: up to 800 KB and 2,000 records. Split larger sheets. Keep this master workbook as XLSX.'],
 ['Evidence','Use exact HTTPS sourceUrl evidence and verificationBasis. Template/example sources are not product evidence.'],
 ['Blank rows','Completely empty data rows are ignored by the importer. Export only populated data sheets. Do not export Instructions.'],
 ['Existing CSVs','For an existing CSV use Excel Data → From Text/CSV and keep identifiers/part numbers/schedule as Text. Double-clicking CSVs can remove leading zeros.'],
 ['More specifications','The wide material-template.csv lists all supported fields. Additional defining text attributes can use attribute.fieldName headers; exclude brand/vendor/part metadata.']]
 for row in notes:guide.append(row)
 guide.column_dimensions['A'].width=26;guide.column_dimensions['B'].width=110;guide.freeze_panes='B2';guide.auto_filter.ref=f'A1:B{len(notes)}'
 for row in guide:
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for row in range(1,len(notes)+1):guide.row_dimensions[row].height=42
 for index,(title,kind,category,attrs) in enumerate(sheets):
  ws=wb.create_sheet(title);headers=base+attrs if kind=='material' else columns[kind]
  ws.append(headers)
  if fictional and (kind!='material' or category=='bracket'):ws.append([examples[kind].get(h,'') for h in headers])
  else:ws.append([None]*len(headers))
  ws.freeze_panes='C2';ws.sheet_view.zoomScale=85
  for col,header in enumerate(headers,1):
   letter=ws.cell(1,col).column_letter
   descriptions={'kind':'Record type: material, product, translation or image. Use the same kind throughout a CSV file.','externalKey':'Your stable catalog item ID. Keep it unique and preserve leading zeros.','materialKey':'The externalKey of the catalog item this row belongs to.','productUnit':'Product/stocking unit; no quantity conversions are inferred.','verification':'provisional for AI-prepared data; reviewed only after independent review.','origin':'How this information was obtained. AI imports use ai_import; generated phrases use ai_generated.','sourceUrl':'HTTPS evidence URL for the exact product/specification or phrase facts.','verificationBasis':'Explain which defining facts the source supports and any unresolved uncertainty.','partNumber':'Manufacturer part number, stored as text exactly as supplied.','phrase':'One search phrase per row. Repeat materialKey for additional phrases.','nominalSpacing':'Nominal size/spacing with units; keep separate from actual minSpanIn/maxSpanIn.','rightsBasis':'Document the image usage rights or permission; a public URL is not permission.'}
   ws.cell(1,col).comment=Comment(descriptions.get(header,'Enter decimal inches as a number.' if header in numeric else 'Enter the '+header+' specification as text; keep the header unchanged.'),'EtalonOps')
   ws.column_dimensions[letter].width=60 if header in ['sourceUrl','verificationBasis','rightsBasis','phrase','specification'] else 34 if header in ['title','manufacturer','externalKey','materialKey','partNumber'] else 23
   for row in range(2,502):
    cell=ws.cell(row,col);cell.number_format='0.########' if header in numeric else '@';cell.alignment=Alignment(vertical='top',wrap_text=True)
   choices={'kind':[kind],'category':[category] if category else [],'productUnit':['EA','FT','STICK','BOX','ROLL','PAIR'],'verification':['provisional','reviewed'],'origin':['ai_import','admin_reviewed','ai_discovery'] if kind in ['material','product'] else ['ai_generated','ai_import','admin_reviewed']}
   if header in choices and choices[header]:
    validation=DataValidation(type='list',formula1='"'+','.join(choices[header])+'"',allow_blank=False);validation.errorTitle='Use a listed value';validation.error='Choose one of the listed values.';validation.showErrorMessage=True;validation.errorStyle='stop';ws.add_data_validation(validation);validation.add(f'{letter}2:{letter}2001')
  table=Table(displayName=f'CatalogSheet{index+1}',ref=f'A1:{ws.cell(1,len(headers)).column_letter}2');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
  ws.row_dimensions[1].height=32;ws.row_dimensions[2].height=54
 for ws in wb:
  for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(vertical='center',wrap_text=True)
 path=out/('EtalonOps-catalog-FICTIONAL-EXAMPLE.xlsx' if fictional else 'EtalonOps-catalog-template.xlsx');wb.save(path)
 with path.open('rb') as f:check=load_workbook(f,data_only=False)
 assert len(check.sheetnames)==8 and check['Manufacturer Products']['D2'].number_format=='@'
 assert not any(cell.data_type=='f' for ws in check for row in ws for cell in row)
 print(f'Generated {path.name}')
